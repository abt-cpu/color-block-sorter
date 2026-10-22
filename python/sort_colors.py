#!/usr/bin/env python3
"""把 Word 檔案裡「螢光筆底色相同」的文字按原順序集中在一起。

  python3 sort_colors.py 檔案.docx [更多.docx ...]    -> 檔案_分類.docx
  python3 sort_colors.py 檔案.docx --report           -> 另存 檔案_驗證報告.md
  python3 sort_colors.py --selftest                   -> 發給別人前先自測
  python3 sort_colors.py                              -> 跳出選檔視窗

內建 5 項檢測；任何一項不過，就不產出分類檔，只留驗證報告。
"""
import os
import random
import sys
from collections import OrderedDict

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

NAMES = {
    "yellow": "黃", "green": "亮綠", "cyan": "青", "magenta": "粉", "blue": "藍", "red": "紅",
    "darkBlue": "深藍", "darkCyan": "深青", "darkGreen": "深綠", "darkMagenta": "紫",
    "darkRed": "深紅", "darkYellow": "深黃", "darkGray": "深灰", "lightGray": "淺灰",
    "black": "黑", "white": "白", "NONE": "無底色",
}
W_T, W_TAB, W_BR, W_CR = qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")


# ---------- 讀取 ----------
def run_text(r):
    out = []
    for c in r:
        if c.tag == W_T:
            out.append(c.text or "")
        elif c.tag == W_TAB:
            out.append("\t")
        elif c.tag in (W_BR, W_CR):
            out.append("\n")
    return "".join(out)


def read_segments(path):
    """依文件順序讀出 (編號, 段落序號, 顏色, 文字)；另回傳「另一條路徑」讀到的總字數。"""
    doc = Document(path)
    segs, sid, total = [], 0, 0
    for pi, p in enumerate(doc.element.body.iter(qn("w:p"))):
        total += len(Paragraph(p, None).text)
        for r in p.xpath("./w:r | ./w:hyperlink/w:r"):
            text = run_text(r)
            if not text:
                continue
            hl = r.find(qn("w:rPr") + "/" + qn("w:highlight"))
            color = hl.get(qn("w:val")) if hl is not None else "NONE"
            segs.append((sid, pi, "NONE" if color == "none" else color, text))
            sid += 1
    return segs, total


def group(segs):
    """{顏色: [段落[(編號, 文字), ...], ...]}，同段落連續同色合併；無底色放最後。"""
    g, last = OrderedDict(), {}
    for sid, pi, c, t in segs:
        paras = g.setdefault(c, [])
        if last.get(c) == pi:
            paras[-1].append((sid, t))
        else:
            paras.append([(sid, t)])
        last[c] = pi
    if "NONE" in g:
        g["NONE"] = g.pop("NONE")
    return g


def count(paras):
    return sum(len(t) for p in paras for _, t in p)


# ---------- 輸出 ----------
def make_run(text, color):
    """直接組 w:r，避開 python-docx 逐次插入時的線性搜尋（大檔案會變成平方時間）。"""
    r = OxmlElement("w:r")
    if color != "NONE":
        rpr = OxmlElement("w:rPr")
        hl = OxmlElement("w:highlight")
        hl.set(qn("w:val"), color)
        rpr.append(hl)
        r.append(rpr)
    for i, line in enumerate(text.split("\n")):
        if i:
            r.append(OxmlElement("w:br"))
        for j, piece in enumerate(line.split("\t")):
            if j:
                r.append(OxmlElement("w:tab"))
            if piece:
                t = OxmlElement("w:t")
                t.text = piece
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r.append(t)
    return r


def write_docx(g, total, path):
    out = Document()
    body = out.element.body
    sect = body.find(qn("w:sectPr"))

    def add_p(runs=()):
        p = OxmlElement("w:p")
        for r in runs:
            p.append(r)
        sect.addprevious(p)

    note = out.add_paragraph(f"原文 {total} 字，已按底色分組並通過 5 項完整性檢測。")
    note.runs[0].font.italic = True
    for c, paras in g.items():
        out.add_heading(f"{NAMES.get(c, c)}　{count(paras)} 字", level=2)
        for para in paras:
            add_p(make_run(t, c) for _, t in para)
    out.save(path)


# ---------- 檢測 ----------
def check_memory(segs, total, g):
    items = [(i, t) for paras in g.values() for para in paras for i, t in para]
    return [
        ("字數守恆（片段總和 = 原文總字數）", sum(len(s[3]) for s in segs) == total,
         f"{sum(len(s[3]) for s in segs)} vs {total}"),
        ("編號不重複不遺漏", sorted(i for i, _ in items) == list(range(len(segs))), f"{len(items)}/{len(segs)}"),
        ("還原測試（排回原序後逐字一致）",
         "".join(t for _, t in sorted(items)) == "".join(s[3] for s in segs), ""),
    ]


def check_file(g, path):
    d = Document(path)
    groups, cur, bad = [], None, 0
    for p in d.paragraphs:
        if (p._p.style or "").startswith("Heading"):
            groups.append([])
            cur = list(g)[len(groups) - 1] if len(groups) <= len(g) else None
        elif groups:
            groups[-1].append(p.text)
            for r in p.runs:
                hl = r.font._element.rPr.find(qn("w:highlight")) if r.font._element.rPr is not None else None
                if (hl.get(qn("w:val")) if hl is not None else "NONE") != cur:
                    bad += 1
    exp = [["".join(t for _, t in para) for para in paras] for paras in g.values()]
    return [
        ("輸出檔重讀後各組內容一致", groups == exp, ""),
        ("輸出檔每個字的底色與其所在組相符", bad == 0, f"異常 run 數 {bad}"),
    ]


# ---------- 主流程 ----------
def report_text(src, segs, total, g, checks, warns):
    L = [f"# 驗證報告：{os.path.basename(src)}", "", "## 各組統計"]
    L += [f"- {NAMES.get(c, c)}（{c}）：{count(p)} 字 / {len(p)} 段" for c, p in g.items()]
    L += [f"- 總計：{total} 字 / {len(segs)} 個片段", "", "## 檢測結果"]
    L += [f"- {'PASS' if ok else 'FAIL'}　{n} {d}" for n, ok, d in checks]
    L += ["", f"## 需人工複查（{len(warns)}）"] + [f"- {w}" for w in warns[:200]]
    return "\n".join(L)


def process(src, want_report=False):
    base = os.path.splitext(src)[0]
    dst, tmp = base + "_分類.docx", base + "_分類.tmp.docx"
    segs, total = read_segments(src)
    if not segs:
        print(f"❌ {src}：讀不到任何文字")
        return False
    g = group(segs)
    checks = check_memory(segs, total, g)
    write_docx(g, total, tmp)
    checks += check_file(g, tmp)
    ok = all(c[1] for c in checks)

    warns, by_par = [], {}
    for sid, pi, c, t in segs:
        by_par.setdefault(pi, []).append((c, t))
    for pi, items in by_par.items():
        if len({c for c, _ in items}) > 1:
            warns.append(f"段落#{pi} 句中換色：" + " | ".join(f"{NAMES.get(c, c)}:{t[:20]}" for c, t in items))
    warns += [f"片段#{s} 是純空白但有底色({c})" for s, _, c, t in segs if c != "NONE" and not t.strip()]

    if ok:
        os.replace(tmp, dst)
    elif os.path.exists(tmp):
        os.remove(tmp)
    if want_report or not ok:
        open(base + "_驗證報告.md", "w", encoding="utf-8").write(report_text(src, segs, total, g, checks, warns))

    sizes = "　".join(f"{NAMES.get(c, c)} {count(p)}" for c, p in g.items())
    print(f"{'✅' if ok else '❌'} {os.path.basename(src)}　共 {total} 字　{sizes}")
    if ok:
        print(f"   → {os.path.basename(dst)}　（5 項檢測全過）")
    else:
        print("   ✗ 檢測未過，沒有產出分類檔：" + "；".join(n for n, k, _ in checks if not k))
    none_share = count(g["NONE"]) / max(total, 1) if "NONE" in g else 0
    if none_share > 0.5:
        print(f"   ⚠ 無底色佔 {none_share:.0%}：若你以為已塗色，檔案可能用的不是螢光筆（而是底紋/字型色）")
    return ok


# ---------- 自測 ----------
SAMPLE = ("這是一段用來測試的文字，包含中文、English words、數字 12345，以及標點符號！\n"
          "第二段：句子會在任何位置被換色，甚至在單詞 mid-dle 的中間。\tTab 也要保留。\n\n"
          "Third paragraph with 空白 and  double  spaces.\n") * 40


def selftest():
    import tempfile
    from docx.enum.text import WD_COLOR_INDEX as C
    cols = [C.YELLOW, C.BRIGHT_GREEN, C.PINK, C.TURQUOISE, C.RED, C.BLUE, None]
    allok = True
    with tempfile.TemporaryDirectory() as d:
        for seed in (1, 2, 3):
            rnd, doc, truth = random.Random(seed), Document(), {}
            for para in [p for p in SAMPLE.split("\n") if p]:
                p, i = doc.add_paragraph(), 0
                while i < len(para):
                    seg = para[i:i + rnd.choice([1, 3, 8, 20, 60])]
                    i += len(seg)
                    c = rnd.choice(cols)
                    r = p.add_run(seg)
                    if c is not None:
                        r.font.highlight_color = c
                    truth[c.name if c else "NONE"] = truth.get(c.name if c else "NONE", "") + seg
            src = os.path.join(d, f"t{seed}.docx")
            doc.save(src)
            ok = process(src)
            # 與「生成時記下的標準答案」獨立比對（用 python-docx 官方介面讀輸出檔，不走本程式的讀取路徑）
            got = {}
            for p in Document(os.path.join(d, f"t{seed}_分類.docx")).paragraphs:
                for r in p.runs:
                    k = r.font.highlight_color.name if r.font.highlight_color else "NONE"
                    if not (p._p.style or "").startswith("Heading") and not p.text.startswith("原文 "):
                        got[k] = got.get(k, "") + r.text
            same = got == truth
            print(("   ✅" if same else "   ❌") + " 與生成時的標準答案逐字一致")
            allok &= ok and same
        # 故意弄壞：刪一個字，檢測必須報錯
        segs, total = read_segments(src)
        g = group(segs)
        i, t = g[next(iter(g))][0][0]
        g[next(iter(g))][0][0] = (i, t[:-1])
        caught = not all(ok for _, ok, _ in check_memory(segs, total, g))
        print("✅ 故意刪掉一個字，檢測有抓到" if caught else "❌ 檢測沒抓到被刪的字！")
        allok &= caught
    print("\n自測" + ("全部通過，可以發給別人使用。" if allok else "失敗，請勿釋出。"))
    return allok


def show(text, ok):
    """沒有終端機時（雙擊執行檔）用視窗顯示結果；設了 SORT_COLORS_LOG 則改寫入該檔案（供自動化測試）。"""
    log = os.environ.get("SORT_COLORS_LOG")
    if log:
        open(log, "a", encoding="utf-8").write(text + "\n")
        return
    from tkinter import messagebox
    (messagebox.showinfo if ok else messagebox.showerror)("螢光底色分類器", text)


def run_quiet(files):
    """逐檔處理並收集輸出，不依賴 stdout（視窗版執行檔的 stdout 可能是 None）。"""
    import contextlib
    import io
    buf, ok = io.StringIO(), True
    with contextlib.redirect_stdout(buf):
        for f in files:
            ok &= process(f)
    show(buf.getvalue().strip(), ok)
    return 0 if ok else 1


def pick_files():
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    files = filedialog.askopenfilenames(title="選擇要分類的 Word 檔（可多選）", filetypes=[("Word 檔", "*.docx")])
    root.destroy()
    return list(files)


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):      # Windows 主控台常是 cp1252/cp950：不能編碼的符號換成 ?，不要當掉
        if stream is not None and hasattr(stream, "reconfigure"):
            stream.reconfigure(errors="replace")
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if "--selftest" in sys.argv:
        sys.exit(0 if selftest() else 1)
    if "--help" in sys.argv or "-h" in sys.argv:
        print(__doc__)
        sys.exit(0)
    if not args:                       # 雙擊執行：跳出選檔視窗
        try:
            args = pick_files()
        except Exception:              # 沒有圖形介面時退回說明
            print(__doc__)
            sys.exit(2)
        if not args:
            sys.exit(0)
        sys.exit(run_quiet(args))
    if sys.stdout is None or os.environ.get("SORT_COLORS_LOG"):   # 視窗版或自動化測試
        sys.exit(run_quiet(args))
    results = [process(a, "--report" in sys.argv) for a in args]
    sys.exit(0 if all(results) else 1)
