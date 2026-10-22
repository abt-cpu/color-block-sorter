"""產生測試用 .docx（固定種子，含 tab、換行、空白、emoji、表格、多種顏色）。
用法: python3 make_samples.py 輸出資料夾
"""
import random
import sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_COLOR_INDEX as C

out = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
out.mkdir(parents=True, exist_ok=True)
COLS = [C.YELLOW, C.BRIGHT_GREEN, C.PINK, C.TURQUOISE, C.RED, C.BLUE, C.DARK_BLUE, None]


def make(path, text, seed, table=False):
    rnd, doc = random.Random(seed), Document()
    for para in [p for p in text.split("\n") if p]:
        p, chars, i = doc.add_paragraph(), list(para), 0
        while i < len(chars):
            n = rnd.choice([1, 2, 5, 12, 40])
            r = p.add_run("".join(chars[i:i + n]))
            i += n
            c = rnd.choice(COLS)
            if c is not None:
                r.font.highlight_color = c
    if table:
        t = doc.add_table(rows=2, cols=2)
        for row in t.rows:
            for cell in row.cells:
                r = cell.paragraphs[0].add_run("表格內")
                r.font.highlight_color = C.YELLOW
                cell.paragraphs[0].add_run("無色")
    doc.save(path)


base = "中文與English混合😀，tab\t分隔\n換行段  兩個空格。\n\n"
make(out / "mixed.docx", base * 60, 11, table=True)
make(out / "large.docx", base * 3000, 12)
print("ok:", sorted(p.name for p in out.glob("*.docx")))
